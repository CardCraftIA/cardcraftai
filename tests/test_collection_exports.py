import unittest
from io import BytesIO
from collection_exports import export_collection
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.items = [dict(card_name='Pokémon & <Pikachu>', set_name='Promo',
                           card_number='001', quantity=2, condition='Near Mint',
                           language='Português', variant='=1+1', wishlist=False,
                           notes='PRIVATE_SENTINEL', user_id='OWNER_SECRET')]

    def text(self, fmt, notes=False):
        data = BytesIO(export_collection(self.items, fmt, notes))
        if fmt == 'pdf':
            return '\n'.join(p.extract_text() for p in PdfReader(data).pages)
        if fmt == 'docx':
            return '\n'.join(p.text for p in Document(data).paragraphs)
        sheet = load_workbook(data).active
        return '\n'.join(str(c.value) for row in sheet for c in row)

    def test_private_fields_excluded_by_default(self):
        for fmt in ('xlsx', 'docx', 'pdf'):
            with self.subTest(fmt=fmt):
                text = self.text(fmt)
                self.assertIn('Pikachu', text)
                self.assertNotIn('PRIVATE_SENTINEL', text)
                self.assertNotIn('OWNER_SECRET', text)

    def test_notes_are_explicitly_opt_in(self):
        for fmt in ('xlsx', 'docx', 'pdf'):
            with self.subTest(fmt=fmt):
                self.assertIn('PRIVATE_SENTINEL', self.text(fmt, True))

    def test_excel_preserves_numbers_and_neutralizes_formulas(self):
        sheet = load_workbook(BytesIO(export_collection(self.items, 'xlsx'))).active
        self.assertEqual(sheet['C2'].value, '001')
        self.assertEqual(sheet['D2'].value, 2)
        self.assertEqual(sheet['G2'].value, '=1+1')
        self.assertEqual(sheet['G2'].data_type, 's')

    def test_empty_and_multi_page_exports(self):
        for fmt in ('xlsx', 'docx', 'pdf'):
            self.assertTrue(export_collection([], fmt))
        self.assertGreater(len(PdfReader(BytesIO(export_collection(self.items * 30, 'pdf'))).pages), 1)


if __name__ == '__main__':
    unittest.main()
