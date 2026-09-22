"""On-demand, in-memory exports of the owner's filtered collection."""
from io import BytesIO
from xml.sax.saxutils import escape

FIELDS = ('card_name', 'set_name', 'card_number', 'quantity', 'condition', 'language', 'variant', 'wishlist')
LABELS = ('Carta / Card', 'Coleção / Set', 'Número / Number', 'Quantidade / Quantity', 'Condição / Condition', 'Idioma / Language', 'Variante / Variant', 'Desejos / Wishlist')


def export_collection(items, format_name, include_notes=False):
    """No network access, no formulas, no hidden private fields. Returns bytes."""
    fields = FIELDS + (('notes',) if include_notes else ())
    labels = LABELS + (('Notas privadas / Private notes',) if include_notes else ())
    output = BytesIO()
    if format_name == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        book = Workbook()
        sheet = book.active
        sheet.title = 'CardCraftAI'
        sheet.append(labels)
        for item in items:
            sheet.append([item.get(key, '') for key in fields])
            for cell in sheet[sheet.max_row]:
                # User text must never become an Excel formula.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                    cell.number_format = '@'
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='123C37')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[1].height = 32
        for index in range(1, len(fields) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 28
        if include_notes:
            sheet.column_dimensions[get_column_letter(len(fields))].width = 64
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        book.save(output)
    elif format_name == 'docx':
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        doc.styles['Normal'].font.size = Pt(10)
        doc.add_heading('CardCraftAI — Coleção / Collection', 0)
        doc.add_paragraph(f'{len(items)} registros / entries')
        for item in items:
            doc.add_heading(str(item.get('card_name', '')), 2)
            for field, label in zip(fields[1:], labels[1:]):
                doc.add_paragraph(f'{label}: {item.get(field, "")}')
        doc.save(output)
    elif format_name == 'pdf':
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        styles = getSampleStyleSheet()
        styles['Heading1'].textColor = colors.HexColor('#123C37')
        story = [Paragraph('CardCraftAI', styles['Title']), Paragraph('Coleção / Collection', styles['Heading1']), Spacer(1, 12)]
        for item in items:
            story.append(Paragraph(escape(str(item.get('card_name', ''))), styles['Heading2']))
            for field, label in zip(fields[1:], labels[1:]):
                text = escape(f'{label}: {item.get(field, "")}').replace('\n', '<br/>')
                story.append(Paragraph(text, styles['BodyText']))
            story.append(Spacer(1, 14))
        SimpleDocTemplate(output, title='CardCraftAI collection').build(story)
    else:
        raise ValueError('Unsupported export format')
    return output.getvalue()
